from __future__ import annotations

from hashlib import sha256
from io import BytesIO
import json
from zipfile import ZIP_DEFLATED, ZipFile

import fitz
from openpyxl import Workbook
import pytest

from service.auth import (
    GrantStatus,
    ServiceAccountGrant,
    WorkerPool,
    resolve_service_account_actor,
)
from service.object_store import InMemoryObjectStore
from service.processing import ArtifactManifest, ClaimedStepAttempt
from service.processing.document_worker import (
    DOCUMENT_PARSER_PROFILE_VERSION,
    DocumentWorkerService,
    InMemoryDocumentRepository,
    build_document_step_definitions,
    document_step_handlers,
)
from service.processing.parsers import (
    ParserQualityError,
    ParserRegistry,
    SourceDocument,
)


def _hash(value: bytes) -> str:
    return sha256(value).hexdigest()


def _docx_bytes() -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            (
                '<?xml version="1.0"?>'
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Override PartName="/word/document.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument.'
                'wordprocessingml.document.main+xml"/>'
                "</Types>"
            ),
        )
        archive.writestr(
            "word/document.xml",
            (
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                "<w:body><w:p><w:r><w:t>Primary endpoint is PFS.</w:t></w:r></w:p>"
                "<w:tbl><w:tr><w:tc><w:p><w:r><w:t>PARAMCD</w:t></w:r></w:p></w:tc>"
                "<w:tc><w:p><w:r><w:t>PFS</w:t></w:r></w:p></w:tc></w:tr></w:tbl>"
                "</w:body></w:document>"
            ),
        )
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CT"
    sheet.append(["Code", "Decode"])
    sheet.append(["Y", "Yes"])
    sheet.append(["TOTAL", "=COUNTA(A2:A2)"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(text: str) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    result = document.tobytes()
    document.close()
    return result


@pytest.mark.parametrize(
    ("media_type", "content", "branch", "locator_kind"),
    [
        ("text/plain", b"Treatment-emergent adverse event.", "text", "line_range"),
        ("text/markdown", b"# Endpoint\n\nPFS definition.", "text", "line_range"),
        (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            _docx_bytes(),
            "text",
            "paragraph",
        ),
        (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            _docx_bytes(),
            "tables",
            "table",
        ),
        (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            _xlsx_bytes(),
            "tables",
            "cell_range",
        ),
        ("application/pdf", _pdf_bytes("PFS is the primary endpoint."), "text", "page"),
    ],
)
def test_deterministic_parsers_emit_stable_locator_and_source_provenance(
    media_type: str,
    content: bytes,
    branch: str,
    locator_kind: str,
) -> None:
    source = SourceDocument(
        source_id="src-test",
        source_version_id="srcv-test",
        source_sha256=_hash(content),
        media_type=media_type,
        object_key="sources/src-test/srcv-test/original.bin",
        data_boundary="local_processing_only",
        rights={"classification": "internal", "storage_allowed": True},
    )
    parser = ParserRegistry.default().for_media_type(media_type)

    result = parser.parse(source=source, content=content, branch=branch)

    assert result.source_version_id == source.source_version_id
    assert result.source_sha256 == source.source_sha256
    assert result.parser_profile_version == DOCUMENT_PARSER_PROFILE_VERSION
    assert result.branch == branch
    assert result.fragments
    assert all(fragment.locator["kind"] == locator_kind for fragment in result.fragments)
    assert all(
        fragment.content_sha256 == _hash(fragment.content.encode()) for fragment in result.fragments
    )
    assert result == parser.parse(source=source, content=content, branch=branch)


def test_scanned_pdf_without_text_fails_closed_instead_of_inventing_ocr() -> None:
    document = fitz.open()
    document.new_page()
    content = document.tobytes()
    document.close()
    source = SourceDocument(
        source_id="src-scan",
        source_version_id="srcv-scan",
        source_sha256=_hash(content),
        media_type="application/pdf",
        object_key="sources/src-scan/srcv-scan/original.pdf",
        data_boundary="local_processing_only",
        rights={"classification": "internal", "storage_allowed": True},
    )

    with pytest.raises(ParserQualityError, match="OCR"):
        ParserRegistry.default().for_media_type("application/pdf").parse(
            source=source,
            content=content,
            branch="text",
        )


def test_xlsx_parser_preserves_formula_text_for_traceability() -> None:
    content = _xlsx_bytes()
    source = SourceDocument(
        source_id="src-formula",
        source_version_id="srcv-formula",
        source_sha256=_hash(content),
        media_type=("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        object_key="sources/src-formula/srcv-formula/original.xlsx",
        data_boundary="local_processing_only",
        rights={"classification": "internal", "storage_allowed": True},
    )

    result = (
        ParserRegistry.default()
        .for_media_type(source.media_type)
        .parse(
            source=source,
            content=content,
            branch="tables",
        )
    )

    assert any("=COUNTA(A2:A2)" in fragment.content for fragment in result.fragments)


def test_document_dag_declares_conditional_branches_and_fan_in() -> None:
    pdf_steps = build_document_step_definitions(
        media_type="application/pdf",
        input_sha256="a" * 64,
    )
    markdown_steps = build_document_step_definitions(
        media_type="text/markdown",
        input_sha256="b" * 64,
    )

    assert [step.step_key for step in pdf_steps] == [
        "document.validate",
        "document.parse_text",
        "document.parse_tables",
        "document.parse_images",
        "document.persist_evidence",
    ]
    assert pdf_steps[-1].depends_on == (
        "document.parse_text",
        "document.parse_tables",
        "document.parse_images",
    )
    assert [step.step_key for step in markdown_steps] == [
        "document.validate",
        "document.parse_text",
        "document.persist_evidence",
    ]


def test_document_handlers_write_derived_artifact_then_evidence_only_at_fan_in() -> None:
    content = b"# Analysis\n\nSafety population includes treated subjects."
    store = InMemoryObjectStore()
    original = store.put_bytes(
        "sources/src-test/srcv-test/original.md",
        content,
        media_type="text/markdown",
    )
    repository = InMemoryDocumentRepository(
        source=SourceDocument(
            source_id="src-test",
            source_version_id="srcv-test",
            source_sha256=original.sha256,
            media_type=original.media_type,
            object_key=original.object_key,
            data_boundary="local_processing_only",
            rights={"classification": "internal", "storage_allowed": True},
        )
    )
    service = DocumentWorkerService(
        repository=repository,
        object_store=store,
        parsers=ParserRegistry.default(),
    )
    handlers = document_step_handlers(service)

    validate = handlers["document.validate"](_context("document.validate", 1))
    parsed = handlers["document.parse_text"](_context("document.parse_text", 2))
    repository.record_successful_manifest(
        "run-test",
        "document.parse_text",
        parsed.artifact_manifest,
    )

    assert validate.artifact_manifest == ArtifactManifest()
    assert repository.evidence == []
    derived = parsed.artifact_manifest.artifacts[0]
    parsed_payload = json.loads(store.get_bytes(derived.object_key))
    assert parsed_payload["parser_profile_version"] == DOCUMENT_PARSER_PROFILE_VERSION
    assert parsed_payload["source_sha256"] == original.sha256

    persisted = handlers["document.persist_evidence"](_context("document.persist_evidence", 3))

    assert persisted.output_sha256
    assert len(repository.evidence) == 1
    evidence = repository.evidence[0]
    assert evidence["source_artifact_kind"] == "original"
    assert evidence["derived_artifact_kind"] == "parser_output"
    assert evidence["parser_profile_version"] == DOCUMENT_PARSER_PROFILE_VERSION
    assert repository.run_status == "author_confirmation_required"
    assert repository.candidates == []
    assert repository.releases == []


def test_parse_retry_reuses_committed_derived_object_without_deleting_it() -> None:
    content = b"# Analysis\n\nSafety population includes treated subjects."
    store = InMemoryObjectStore()
    original = store.put_bytes(
        "sources/src-test/srcv-test/original.md",
        content,
        media_type="text/markdown",
    )
    repository = InMemoryDocumentRepository(
        source=SourceDocument(
            source_id="src-test",
            source_version_id="srcv-test",
            source_sha256=original.sha256,
            media_type=original.media_type,
            object_key=original.object_key,
            data_boundary="local_processing_only",
            rights={"classification": "internal", "storage_allowed": True},
        )
    )
    service = DocumentWorkerService(
        repository=repository,
        object_store=store,
        parsers=ParserRegistry.default(),
    )

    first = service.parse(_context("document.parse_text", 10), branch="text")
    retried = service.parse(_context("document.parse_text", 11), branch="text")

    assert retried == first
    derived = first.artifact_manifest.artifacts[0]
    assert store.head(derived.object_key) == derived
    assert len(repository.derived) == 1
    assert set(repository.write_intents.values()) == {"committed"}


def _context(step_key: str, ordinal: int):
    claim = ClaimedStepAttempt(
        run_id="run-test",
        step_id=f"step-{ordinal}",
        step_key=step_key,
        pool=WorkerPool.DOCUMENT,
        attempt_id=f"attempt-{ordinal}",
        attempt_number=1,
        input_sha256="a" * 64,
    )
    return type(
        "Context",
        (),
        {
            "claim": claim,
            "heartbeat": lambda self: None,
            "checkpoint": lambda self, value: None,
        },
    )()


def test_document_worker_actor_never_receives_governance_or_release_permissions() -> None:
    actor = resolve_service_account_actor(
        ServiceAccountGrant(
            service_account_id="svc-document",
            display_name="Document Worker",
            worker_pool=WorkerPool.DOCUMENT,
            scopes=[
                "source:read",
                "object:read",
                "object:write_derived",
                "processing:execute",
                "evidence:write",
            ],
            secret_ref="env://P12_DOCUMENT_WORKER_TOKEN",
            status=GrantStatus.ACTIVE,
        )
    )

    assert {
        "candidate:write",
        "review:decide",
        "release:build",
        "release:publish",
        "index:build",
    }.isdisjoint(permission.value for permission in actor.permissions)
